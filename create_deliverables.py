from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path
import json, textwrap

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / 'docs'
DOCS.mkdir(exist_ok=True)
with (ROOT/'models'/'evaluation_metrics.json').open() as f:
    M = json.load(f)

NAVY = '171B3A'
PURPLE = '5546A7'
VIOLET = '7A69D6'
CYAN = '4AAAD7'
MINT = '55B89F'
ORANGE = 'F4AE59'
PALE = 'F5F6FB'
WHITE = 'FFFFFF'
TEXT = '202744'
MUTED = '68708B'
BORDER = 'E2E6F0'
RED = 'C85A54'
FONT = 'Aptos'


def rgb(hexv):
    return RGBColor.from_string(hexv)


def setup_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_bg(slide, color=PALE):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, size=18, color=TEXT, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, x, y, w, h, lines, size=15, color=TEXT, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.06); tf.margin_bottom = Inches(0.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT; p.font.size = Pt(size); p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
        p.level = 0
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=True, line_width=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line); sh.line.width = Pt(line_width)
    if radius:
        try: sh.adjustments[0] = 0.12
        except Exception: pass
    return sh


def ellipse(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line or fill)
    return sh


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.5, arrow=True):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color); sh.line.width = Pt(width)
    # Direction is explained by left-to-right layout and labels; use compatible connectors.
    return sh


def tag(slide, x, y, w, text, fill=PURPLE, color=WHITE, size=10):
    rect(slide, x, y, w, 0.33, fill=fill, line=fill, radius=True)
    add_text(slide, x+0.05, y+0.015, w-0.1, 0.28, text, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def footer(slide, num, dark=False):
    color = 'D7DAEC' if dark else MUTED
    add_text(slide, 0.5, 7.15, 8.8, 0.18, 'GLUCOTWIN  •  SYNTHETIC DATA ONLY  •  NOT FOR CLINICAL USE', size=8, color=color, bold=True, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, 12.25, 7.12, 0.55, 0.2, f'{num:02d}', size=9, color=color, bold=True, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE, margin=0)


def slide_title(slide, eyebrow, title, subtitle=None, num=None):
    add_text(slide, 0.65, 0.36, 11.9, 0.25, eyebrow.upper(), size=10, color=PURPLE, bold=True, margin=0)
    add_text(slide, 0.65, 0.72, 12.0, 0.52, title, size=27, color=NAVY, bold=True, margin=0)
    if subtitle:
        add_text(slide, 0.65, 1.34, 12.0, 0.45, subtitle, size=12, color=MUTED, margin=0)
    if num is not None: footer(slide, num)

# ---------------- Pitch presentation ----------------
prs = setup_presentation()

# Slide 1: title
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s, NAVY)
# subtle decorative orbs
ellipse(s, 9.8, -0.8, 4.3, 4.3, '242950')
ellipse(s, 11.1, 3.4, 2.7, 2.7, '22264B')
tag(s, 0.75, 0.72, 2.1, 'DIGITAL TWIN PoC', fill='3B3A75')
add_text(s, 0.75, 1.55, 7.4, 0.95, 'GlucoTwin', size=42, color=WHITE, bold=True, margin=0)
add_text(s, 0.78, 2.58, 7.4, 0.75, 'Predicting a synthetic glucose spike\ntwo hours before it happens', size=23, color='DDE1FF', bold=True, margin=0)
add_text(s, 0.78, 4.05, 6.5, 0.68, 'Fusing historical EHR context with hourly wearable and CGM signals in a virtual patient dashboard.', size=15, color='C2C8E7', margin=0)
# flow visual
cards = [(8.65, 1.3, 3.55, 1.12, '01  EHR PROFILE', 'HbA1c · BMI · history', CYAN),
         (8.65, 2.8, 3.55, 1.12, '02  LIVE-LIKE SENSORS', 'CGM · HRV · sleep · steps', MINT),
         (8.65, 4.3, 3.55, 1.12, '03  2-HOUR FORECAST', 'Model score + what-if UI', ORANGE)]
for x,y,w,h,head,sub,c in cards:
    rect(s,x,y,w,h,fill='252A50',line='454B7A')
    ellipse(s,x+0.22,y+0.3,0.22,0.22,c)
    add_text(s,x+0.58,y+0.18,w-0.8,0.28,head,size=12,color=WHITE,bold=True,margin=0)
    add_text(s,x+0.58,y+0.52,w-0.8,0.35,sub,size=11,color='CDD2EE',margin=0)
line(s,10.42,2.44,10.42,2.77,color='7F88B5',width=1.8)
line(s,10.42,3.94,10.42,4.27,color='7F88B5',width=1.8)
add_text(s, 0.78, 6.15, 6.7, 0.32, 'Happiest Health Digital Twin Challenge 2026', size=12, color='ABB3D8', bold=True, margin=0)
add_text(s, 0.78, 6.55, 7.4, 0.28, 'Research prototype • Synthetic data only • Not for clinical use', size=10, color='9AA3CC', margin=0)
footer(s,1,dark=True)

# Slide 2: problem/use case
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'01 · Problem & use case','A narrow prediction task beats a vague “whole-body twin”','Scope the digital twin around one localized health outcome: an elevated synthetic CGM reading two hours ahead.',2)
rect(s,0.68,2.0,5.55,3.9,fill=WHITE,line=BORDER)
tag(s,0.98,2.27,1.45,'THE GAP',fill='E9E6FB',color=PURPLE)
add_text(s,1.0,2.83,4.9,0.55,'A static record is not the current patient state.',size=21,color=NAVY,bold=True,margin=0)
add_rich_lines(s,1.0,3.55,4.8,1.65,[
    '• Historical EHR contains slower-changing context.',
    '• Wearable signals reflect the current state and recent behavior.',
    '• A clinician needs a concise forward-looking signal, not another data silo.'
],size=14,color=TEXT,gap=12)
rect(s,6.55,2.0,6.05,3.9,fill='ECEBFA',line='D7D3F5')
tag(s,6.9,2.27,1.62,'THE PoC',fill=PURPLE)
add_text(s,6.93,2.83,5.2,0.55,'Forecast the two-hour glucose event.',size=21,color=NAVY,bold=True,margin=0)
# target number
add_text(s,6.94,3.55,2.3,0.7,'≥ 180',size=34,color=PURPLE,bold=True,margin=0)
add_text(s,9.05,3.72,2.65,0.36,'mg/dL at t + 2 h',size=14,color=TEXT,bold=True,margin=0)
add_text(s,6.94,4.55,4.9,0.85,'A demo forecast driven by EHR context, CGM trend, meal logs, activity, sleep, HR and HRV.',size=14,color=TEXT,margin=0)
rect(s,0.7,6.15,11.9,0.56,fill='FFF4DD',line='F4DBA0')
add_text(s,0.92,6.28,11.45,0.25,'Important: “≥180 mg/dL” is the synthetic label definition for this exercise, not an independent clinical recommendation.',size=11,color='775817',bold=True,margin=0)

# Slide 3: architecture
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'02 · System design','Two streams become one patient-specific forecast','Every prediction combines a selected EHR profile and a current / recent sensor window; the future reading is used only as a training label.',3)
# nodes
rect(s,0.7,2.15,2.55,1.16,fill=WHITE,line='B8DFF0')
add_text(s,0.88,2.35,2.2,0.27,'STATIC EHR',size=14,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,0.88,2.72,2.2,0.33,'Age · HbA1c · BMI · history',size=10,color=MUTED,align=PP_ALIGN.CENTER,margin=0)
rect(s,0.7,4.05,2.55,1.16,fill=WHITE,line='B9E3D6')
add_text(s,0.88,4.24,2.2,0.27,'DYNAMIC WEARABLE',size=13,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,0.88,4.6,2.2,0.42,'CGM · HR/HRV · steps\nsleep · logged carbs',size=10,color=MUTED,align=PP_ALIGN.CENTER,margin=0)
rect(s,4.0,3.02,2.3,1.25,fill='EAE8FA',line='CDC7F1')
add_text(s,4.18,3.24,1.95,0.33,'FUSION LAYER',size=13,color=PURPLE,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,4.18,3.67,1.95,0.4,'Join by patient + time\nBuild feature row',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
rect(s,7.1,3.02,2.25,1.25,fill='E8F6F3',line='B6DED6')
add_text(s,7.26,3.24,1.95,0.33,'RANDOM FOREST',size=13,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,7.26,3.67,1.95,0.4,'Predict event probability',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
rect(s,10.1,3.02,2.5,1.25,fill='FFF2DE',line='F2D8A3')
add_text(s,10.25,3.24,2.2,0.33,'DOCTOR-FACING UI',size=12,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,10.25,3.67,2.2,0.4,'Probability · trends\nwhat-if controls',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
line(s,3.25,2.72,4.0,3.4,color=CYAN)
line(s,3.25,4.62,4.0,3.85,color=MINT)
line(s,6.3,3.64,7.1,3.64,color=PURPLE)
line(s,9.35,3.64,10.1,3.64,color=ORANGE)
# training target row
rect(s,4.0,5.0,5.35,0.85,fill='F0F1F7',line=BORDER)
add_text(s,4.2,5.2,4.95,0.25,'RETROSPECTIVE LABEL FOR TRAINING ONLY',size=11,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,4.2,5.49,4.95,0.22,'Observed synthetic CGM at t + 2 h ≥ 180 mg/dL',size=10,color=MUTED,align=PP_ALIGN.CENTER,margin=0)
line(s,5.15,4.27,5.15,5.0,color=MUTED,width=1.2,arrow=False)
add_text(s,0.75,6.25,11.8,0.28,'Validation: patient-level holdout (240 train patients / 80 unseen test patients). Target columns are excluded from model inputs.',size=11,color=MUTED,bold=True,margin=0)

# Slide 4: data privacy
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'03 · Data & privacy','A reproducible sandbox with no real patient data','A deterministic generator creates EHR profiles, hourly sensor streams and a separate future-outcome file.',4)
stats=[('320','synthetic patients',CYAN),('53,760','hourly sensor rows',MINT),('53,120','two-hour labels',PURPLE),('7 days','per patient',ORANGE)]
for i,(val,label,c) in enumerate(stats):
    x=0.7+i*3.1
    rect(s,x,2.15,2.75,1.32,fill=WHITE,line=BORDER)
    add_text(s,x+0.18,2.35,2.37,0.48,val,size=26,color=c,bold=True,align=PP_ALIGN.CENTER,margin=0)
    add_text(s,x+0.14,2.94,2.45,0.28,label,size=11,color=MUTED,bold=True,align=PP_ALIGN.CENTER,margin=0)
rect(s,0.7,3.95,5.72,2.05,fill='EAF7F3',line='C6E9DF')
add_text(s,0.98,4.2,5.1,0.35,'STATIC STREAM · EHR',size=13,color='227D69',bold=True,margin=0)
add_rich_lines(s,0.98,4.7,4.95,1.0,['HbA1c, BMI, age and diabetes duration','Comorbidity and simulated medication fields','One profile per fake patient ID'],size=12,color=TEXT,gap=7)
rect(s,6.75,3.95,5.85,2.05,fill='EBF5FB',line='C4E3F0')
add_text(s,7.05,4.2,5.1,0.35,'DYNAMIC STREAM · WEARABLE / CGM',size=13,color='247EA9',bold=True,margin=0)
add_rich_lines(s,7.05,4.7,5.0,1.0,['Hourly CGM, HR, HRV and step count','Sleep duration and simulated food log','Separate label file stores the future observation'],size=12,color=TEXT,gap=7)
add_text(s,0.78,6.34,11.7,0.25,'Privacy by construction: fixed seed, synthetic IDs, no personal records, no outbound APIs, no secrets.',size=12,color=NAVY,bold=True,margin=0)

# Slide 5: model and split
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'04 · Model & evaluation','Evaluation is patient-held-out—not a random row split','A separate set of 80 virtual patients tests how the model behaves on identities it did not train on.',5)
# split visual
rect(s,0.72,2.0,5.5,1.32,fill='EAE8FA',line='CDC7F1')
add_text(s,0.98,2.25,4.95,0.38,'TRAIN · 75% OF PATIENTS',size=16,color=PURPLE,bold=True,margin=0)
add_text(s,0.98,2.74,4.95,0.3,'240 patients · 39,840 hourly windows',size=12,color=TEXT,margin=0)
rect(s,6.8,2.0,5.75,1.32,fill='E8F6F3',line='B6DED6')
add_text(s,7.08,2.25,5.2,0.38,'TEST · 25% UNSEEN PATIENTS',size=16,color='247F69',bold=True,margin=0)
add_text(s,7.08,2.74,5.2,0.3,'80 patients · 13,280 hourly windows',size=12,color=TEXT,margin=0)
# metrics row
metric_items=[('ROC-AUC',f"{M['roc_auc']:.3f}"),('Average precision',f"{M['average_precision']:.3f}"),('Precision @ 0.50',f"{M['precision_at_0_5']:.3f}"),('Recall @ 0.50',f"{M['recall_at_0_5']:.3f}")]
for i,(label,val) in enumerate(metric_items):
    x=0.72+i*3.12
    rect(s,x,3.8,2.8,1.35,fill=WHITE,line=BORDER)
    add_text(s,x+0.1,4.02,2.6,0.27,label,size=10,color=MUTED,bold=True,align=PP_ALIGN.CENTER,margin=0)
    add_text(s,x+0.1,4.38,2.6,0.55,val,size=26,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
rect(s,0.72,5.55,11.95,0.88,fill='FFF4DD',line='F4DBA0')
add_text(s,0.98,5.73,11.35,0.42,'Metrics are from a generated simulator, with 46.0% synthetic event prevalence in the test set. They are not clinical performance estimates.',size=12,color='775817',bold=True,valign=MSO_ANCHOR.MIDDLE,margin=0)

# Slide 6: dashboard walkthrough
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'05 · Clinician-facing concept UI','A compact patient view plus a controlled “what-if” replay','The Streamlit app supports patient selection, timeline replay, trend charts, model score, and sandboxed sensor edits.',6)
# draw browser frame
rect(s,0.68,1.95,11.95,4.5,fill=WHITE,line='CFD5E5')
rect(s,0.68,1.95,11.95,0.47,fill=NAVY,line=NAVY,radius=False)
for i,c in enumerate([RED,ORANGE,MINT]): ellipse(s,0.9+i*0.24,2.1,0.11,0.11,c)
add_text(s,1.83,2.07,7.3,0.22,'GlucoTwin  /  Virtual Patient',size=10,color=WHITE,bold=True,margin=0)
# left panel
rect(s,0.92,2.66,2.15,3.42,fill='F4F5FA',line='E0E4EF')
add_text(s,1.12,2.85,1.75,0.24,'EHR PROFILE',size=11,color=PURPLE,bold=True,margin=0)
for i,(a,b) in enumerate([('Age','57 y'),('HbA1c','8.1%'),('BMI','29.4'),('Duration','9 y'),('Adherence','82%')]):
    yy=3.28+i*0.48
    add_text(s,1.12,yy,0.78,0.21,a,size=9,color=MUTED,margin=0)
    add_text(s,1.9,yy,0.88,0.21,b,size=9,color=TEXT,bold=True,margin=0)
# top metrics
for i,(lab,val,c) in enumerate([('CURRENT CGM','164 mg/dL',PURPLE),('HR / HRV','82 bpm / 41 ms',CYAN),('2-HOUR SCORE','68%',ORANGE)]):
    x=3.32+i*2.96
    rect(s,x,2.66,2.7,0.93,fill='F8F9FC',line=BORDER)
    add_text(s,x+0.12,2.82,2.45,0.2,lab,size=8,color=MUTED,bold=True,align=PP_ALIGN.CENTER,margin=0)
    add_text(s,x+0.12,3.13,2.45,0.3,val,size=15,color=c,bold=True,align=PP_ALIGN.CENTER,margin=0)
# timeline shape chart
rect(s,3.32,3.84,5.58,1.8,fill='FBFCFE',line=BORDER)
add_text(s,3.52,3.98,5.1,0.2,'24-HOUR SYNTHETIC CGM TREND',size=9,color=MUTED,bold=True,margin=0)
# sparkline points
points=[(3.65,5.12),(4.1,4.95),(4.55,4.8),(5.0,4.6),(5.45,4.95),(5.9,4.7),(6.35,4.42),(6.8,4.55),(7.25,4.18),(7.7,4.37),(8.15,4.1),(8.6,4.25)]
for (x1,y1),(x2,y2) in zip(points,points[1:]): line(s,x1,y1,x2,y2,color=PURPLE,width=2.3,arrow=False)
line(s,3.55,4.65,8.68,4.65,color='D9928C',width=1,arrow=False)
add_text(s,8.0,4.43,0.7,0.18,'180',size=8,color=RED,bold=True,margin=0)
# scenario panel
rect(s,9.12,3.84,3.18,2.24,fill='FFF6E9',line='F1DBB2')
add_text(s,9.33,4.03,2.72,0.25,'WHAT-IF SENSOR UPDATE',size=9,color='93671C',bold=True,margin=0)
for i,(label,bar) in enumerate([('Glucose',0.68),('Carbs',0.47),('Steps',0.62),('Sleep',0.72)]):
    yy=4.48+i*0.34
    add_text(s,9.33,yy,0.75,0.17,label,size=8,color=TEXT,margin=0)
    rect(s,10.15,yy+0.02,1.45,0.13,fill='E5E0D6',line='E5E0D6')
    rect(s,10.15,yy+0.02,1.45*bar,0.13,fill=ORANGE,line=ORANGE)
add_text(s,0.92,6.18,11.2,0.18,'Conceptual wireframe: actual Streamlit implementation is included in app/app.py.',size=9,color=MUTED,italic=True,margin=0)

# Slide 7: limits/importance
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s)
slide_title(s,'06 · Findings & limitations','Synthetic metrics are not clinical evidence','The next milestone is independent validation—not simply a higher AUC.',7)
rect(s,0.7,2.0,5.55,3.98,fill=WHITE,line=BORDER)
add_text(s,0.98,2.28,4.95,0.38,'TOP SIMULATED SIGNALS',size=13,color=PURPLE,bold=True,margin=0)
features=list(M['feature_importance'].items())[:6]
maxv=max(v for _,v in features)
for i,(name,val) in enumerate(features):
    yy=2.9+i*0.48
    display=name.replace('_',' ')
    add_text(s,0.98,yy,1.95,0.24,display,size=9,color=TEXT,margin=0)
    rect(s,3.0,yy+0.035,2.55,0.16,fill='EEF0F7',line='EEF0F7')
    rect(s,3.0,yy+0.035,2.55*val/maxv,0.16,fill=PURPLE if i<2 else CYAN,line=PURPLE if i<2 else CYAN)
    add_text(s,5.62,yy,0.42,0.2,f'{val:.2f}',size=8,color=MUTED,align=PP_ALIGN.RIGHT,margin=0)
add_text(s,0.98,5.72,4.95,0.18,'Importance is not causality; time-of-day can encode simulator artifacts.',size=9,color=MUTED,italic=True,margin=0)
rect(s,6.65,2.0,5.95,3.98,fill='FFF4DD',line='F4DBA0')
add_text(s,6.98,2.28,5.2,0.38,'WHAT HAS NOT BEEN DONE',size=13,color='93671C',bold=True,margin=0)
add_rich_lines(s,6.98,2.9,5.0,2.6,[
    '• External validation on independently sourced data',
    '• Probability calibration and threshold selection',
    '• Subgroup / fairness assessment',
    '• Prospective alert utility and clinician workflow study',
    '• Clinical, privacy, governance and regulatory review'
],size=12,color=TEXT,gap=10)

# Slide 8: roadmap and close
s=prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s,NAVY)
add_text(s,0.72,0.54,11.6,0.25,'07 · NEXT STEPS',size=10,color='9BA4D5',bold=True,margin=0)
add_text(s,0.72,0.95,11.5,0.65,'From reproducible PoC to credible research',size=29,color=WHITE,bold=True,margin=0)
roadmap=[('01','Data validity','Use governed, appropriate external datasets and document provenance.',CYAN),
         ('02','Model validity','Calibrate probabilities, test subgroups, quantify uncertainty.',MINT),
         ('03','Clinical workflow','Co-design with clinicians; assess alert burden and usability.',ORANGE),
         ('04','Governance','Security, consent/legal basis, privacy, and safety review.',VIOLET)]
for i,(n,title,desc,c) in enumerate(roadmap):
    x=0.74+i*3.11
    rect(s,x,2.15,2.78,2.25,fill='252A50',line='454B7A')
    ellipse(s,x+0.22,2.42,0.45,0.45,c)
    add_text(s,x+0.22,2.53,0.45,0.18,n,size=10,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
    add_text(s,x+0.22,3.1,2.3,0.32,title,size=15,color=WHITE,bold=True,margin=0)
    add_text(s,x+0.22,3.56,2.3,0.62,desc,size=10,color='CCD2ED',margin=0)
rect(s,0.74,5.15,11.97,0.9,fill='30365F',line='555D8F')
add_text(s,1.02,5.4,11.35,0.36,'GlucoTwin demonstrates EHR + wearable fusion, a reproducible model pipeline and an interactive virtual-patient view.',size=14,color=WHITE,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,0.74,6.45,11.7,0.28,'Thank you  •  All displayed patients and measurements are synthetic  •  Not for clinical use',size=11,color='ABB4D8',align=PP_ALIGN.CENTER,margin=0)
footer(s,8,dark=True)

out = DOCS / 'presentation.pptx'
prs.save(out)

# ---------------- Architecture single-slide deck ----------------
arch = setup_presentation()
s=arch.slides.add_slide(arch.slide_layouts[6]); set_bg(s)
slide_title(s,'GlucoTwin · Architecture','EHR + wearable fusion to 2-hour event forecast','Training and inference are separated: retrospective CGM at t+2h defines the training label but is never a prediction input.',1)
# left input boxes
rect(s,0.6,2.1,2.35,1.25,fill='EAF7FC',line='B8DFF0')
add_text(s,0.78,2.29,1.98,0.3,'STATIC EHR',size=14,color='247EA9',bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,0.78,2.72,1.98,0.42,'Demographics · HbA1c\nBMI · comorbidity · medication',size=9,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
rect(s,0.6,4.0,2.35,1.25,fill='EAF7F3',line='B9E3D6')
add_text(s,0.78,4.19,1.98,0.3,'DYNAMIC STREAM',size=13,color='227D69',bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,0.78,4.62,1.98,0.43,'Hourly CGM · HR · HRV\nsteps · sleep · meal log',size=9,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
# merge
rect(s,3.65,3.03,2.15,1.38,fill='EAE8FA',line='CDC7F1')
add_text(s,3.82,3.23,1.82,0.3,'FUSION',size=14,color=PURPLE,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,3.82,3.66,1.82,0.5,'Join by patient_id\nand timestamp',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
line(s,2.95,2.72,3.65,3.4,color=CYAN)
line(s,2.95,4.62,3.65,4.0,color=MINT)
# inference path
rect(s,6.5,3.03,2.15,1.38,fill='F1F0FB',line='D2CDF1')
add_text(s,6.66,3.23,1.82,0.3,'FEATURE ROW',size=13,color=PURPLE,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,6.66,3.66,1.82,0.5,'Current + past inputs\n(no future leakage)',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
line(s,5.8,3.72,6.5,3.72,color=PURPLE)
rect(s,9.25,3.03,1.72,1.38,fill='E8F6F3',line='B6DED6')
add_text(s,9.37,3.23,1.48,0.3,'RANDOM FOREST',size=11,color='227D69',bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,9.37,3.66,1.48,0.5,'Probability\nscore',size=10,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
line(s,8.65,3.72,9.25,3.72,color=MINT)
rect(s,11.45,3.03,1.35,1.38,fill='FFF2DE',line='F2D8A3')
add_text(s,11.54,3.22,1.17,0.32,'STREAMLIT',size=10,color='93671C',bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,11.54,3.64,1.17,0.5,'Trends +\nwhat-if UI',size=9,color=TEXT,align=PP_ALIGN.CENTER,margin=0)
line(s,10.97,3.72,11.45,3.72,color=ORANGE)
# training loop bottom
rect(s,3.65,5.12,6.32,0.83,fill='F1F2F7',line=BORDER)
add_text(s,3.86,5.3,5.9,0.24,'TRAIN / EVALUATE (OFFLINE)',size=12,color=NAVY,bold=True,align=PP_ALIGN.CENTER,margin=0)
add_text(s,3.86,5.61,5.9,0.2,'Target = synthetic CGM at t + 2 h ≥ 180 mg/dL · patient-level 75/25 split',size=9,color=MUTED,align=PP_ALIGN.CENTER,margin=0)
line(s,4.75,4.41,4.75,5.12,color=MUTED,width=1.15,arrow=False)
line(s,8.85,5.12,7.7,4.41,color=MUTED,width=1.15,arrow=False)
rect(s,0.65,6.25,12.05,0.48,fill='FFF4DD',line='F4DBA0')
add_text(s,0.82,6.36,11.7,0.22,'Synthetic-only research prototype; no real patient records, no external API, no clinical use. Future/target values never enter the model input feature list.',size=10,color='775817',bold=True,align=PP_ALIGN.CENTER,margin=0)
arch.save(DOCS/'architecture.pptx')
print('Created', out, 'and', DOCS/'architecture.pptx')
